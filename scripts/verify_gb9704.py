"""用本机 Word / WPS 实际排版，核对公文预设是否满足 GB/T 9704-2012 的版式要求。

    venv\\Scripts\\python scripts\\verify_gb9704.py

检查项：每页 22 行、每行 28 字、首行缩进 2 字（首行 26 字）、各层次标题字体、
页码 4 号宋体一字线、单页居右双页居左空一字、页码位置（一字线上距版心下边缘约 7 mm）。
需要本机安装 Word 或 WPS（两者都装则各测一次）。
"""

import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document  # noqa: E402

from app.engine import apply_preset  # noqa: E402
from app.presets import BUILTIN_PRESETS, merge_preset  # noqa: E402


def make_sample(path):
    d = Document()
    d.add_heading("关于进一步加强党政机关公文格式规范化管理工作的通知", level=1)
    d.add_paragraph("各市、县人民政府，省政府各部门：")
    d.add_paragraph("为进一步规范公文格式，提高公文质量，现将有关事项通知如下。" * 3)
    d.add_heading("一、总体要求", level=2)
    d.add_paragraph("测试正文内容。" * 40)
    d.add_heading("（一）基本原则", level=3)
    d.add_paragraph("测试正文内容。" * 40)
    d.add_heading("1.具体措施", level=4)
    d.add_paragraph("正" * 726)  # 首行 26 字 + 25 行 × 28 字 = 726 → 恰好 26 行
    for _ in range(6):
        d.add_paragraph("测试正文内容。" * 40)
    d.save(path)


def measure(progid, path):
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    try:
        app = win32com.client.DispatchEx(progid)
    except Exception:
        return None
    app.Visible = False
    try:
        app.DisplayAlerts = 0
    except Exception:
        pass
    res = {}
    try:
        doc = app.Documents.Open(path, False, False)
        app.ActiveWindow.View.Type = 3  # 页面视图，才有行信息
        ps = doc.PageSetup
        res["页面设置"] = f"LinesPage={ps.LinesPage:g} CharsLine={ps.CharsLine:g} 网格模式={ps.LayoutMode}"
        sel = app.Selection
        probes = (("正文", "正正正"), ("一级", "一、"), ("二级", "（一）"), ("三级", "1."))
        for tag, prefix in probes:
            para = next(p for p in doc.Paragraphs if p.Range.Text.startswith(prefix))
            sel.SetRange(para.Range.Start, para.Range.Start)
            counts = []
            for _ in range(3):
                sel.HomeKey(5)
                sel.EndKey(5, 1)
                counts.append(len(sel.Text.strip("\r\n")))
                sel.Collapse(0)
                sel.MoveDown(5, 1)
            res[tag] = f"前3行字数={counts} 字体={para.Range.Font.NameFarEast} {para.Range.Font.Size:g}pt"
        big = next(p for p in doc.Paragraphs if p.Range.Text.startswith("正正正"))
        res["726字段落行数"] = big.Range.ComputeStatistics(1)
        p2, p3 = doc.GoTo(1, 1, 2).Start, doc.GoTo(1, 1, 3).Start
        res["第2页行数"] = doc.Range(p2, p3).ComputeStatistics(1)
        for which, label in ((1, "奇数页"), (3, "偶数页")):
            f = doc.Sections(1).Footers(which)
            pr = f.Range.Paragraphs(1)
            res[f"页码{label}"] = (f"{f.Range.Text.strip()!r} 对齐={pr.Alignment} 左缩进={pr.LeftIndent:.0f}pt "
                                 f"右缩进={pr.RightIndent:.0f}pt {f.Range.Font.NameFarEast} {f.Range.Font.Size:g}pt")
        top_mm = doc.Sections(1).Footers(1).Range.Information(6) / 72 * 25.4
        res["页码行顶部距页顶"] = f"{top_mm:.1f} mm（版心下边缘 262 mm，一字线中心约在其下 7 mm）"
        doc.Close(False)
    finally:
        app.Quit()
    return res


def main():
    gb = next(merge_preset(p) for p in BUILTIN_PRESETS if "党政机关公文" in p["name"])
    tmp = tempfile.gettempdir()
    src, out = os.path.join(tmp, "gb9704_src.docx"), os.path.join(tmp, "gb9704_out.docx")
    make_sample(src)
    res = apply_preset(src, gb, out)
    print("排版：", res["message"] if res["ok"] else res["error"])
    ok_any = False
    for progid in ("Word.Application", "KWPS.Application"):
        r = measure(progid, out)
        if r is None:
            print(f"--- {progid}: 未安装，跳过")
            continue
        ok_any = True
        print(f"--- {progid}")
        for k, v in r.items():
            print(f"    {k}: {v}")
        checks = [
            ("每页 22 行", r["第2页行数"] == 22),
            ("每行 28 字 / 首行 26 字", r["726字段落行数"] == 26 and "[26, 28, 28]" in r["正文"]),
            ("一级黑体 / 二级楷体 / 三级仿宋", "黑体" in r["一级"] and "楷体_GB2312" in r["二级"] and "仿宋_GB2312" in r["三级"]),
            ("页码 4 号宋体一字线", "宋体 14pt" in r["页码奇数页"] and "— " in r["页码奇数页"]),
            ("单页居右、双页居左", "对齐=2" in r["页码奇数页"] and "对齐=0" in r["页码偶数页"]),
        ]
        for name, passed in checks:
            print(f"    {'✔' if passed else '✘'} {name}")
    if not ok_any:
        print("本机没有 Word / WPS，无法进行版式核对。")
    print("输出文件：", out)


if __name__ == "__main__":
    main()
