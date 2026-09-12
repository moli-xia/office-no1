"""Batch and GUI regressions; only temporary documents are used."""
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

# 测试不读写真实的用户设置
os.environ["DOCFMT_SETTINGS"] = os.path.join(tempfile.gettempdir(), "docfmt_test_settings.ini")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from PySide6.QtCore import Qt, QEventLoop, QTimer
from PySide6.QtWidgets import QApplication
from app import fonts
from app.gui import FormatWorker, MainWindow
from app.presets import merge_preset
import shutil

class WorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_collision_and_failure_progress(self):
        with tempfile.TemporaryDirectory() as folder:
            src = Path(folder)/'sample.docx'
            doc = Document(); doc.add_paragraph('测试正文'); doc.save(src)
            existing = Path(folder)/'sample_done.docx'; existing.write_bytes(b'keep')
            bad = Path(folder)/'broken.doc'; bad.write_bytes(b'bad')
            worker = FormatWorker([str(bad), str(src)], merge_preset({}), 'suffix', '', '_done')
            progress, outputs, summaries, done = [], [], [], []
            worker.fileDone.connect(lambda p, ok, msg: done.append((Path(p).name, ok)))
            worker.progressSig.connect(lambda v,t: progress.append(v))
            worker.outputReady.connect(lambda s,o: outputs.append(o))
            worker.finishedAll.connect(summaries.append)
            with patch('app.gui.engine.convert_doc_to_docx', return_value=False):
                worker.run()
            self.assertEqual(progress, [1,2])
            self.assertEqual(existing.read_bytes(), b'keep')
            self.assertEqual(Path(outputs[0]).name, 'sample_done (2).docx')
            self.assertIn('失败 1', summaries[0])
            self.assertEqual(done, [('broken.doc', False), ('sample.docx', True)])
            self.assertEqual(Document(src).paragraphs[0].text, '测试正文')

    def test_ui_run_and_dedup(self):
        with tempfile.TemporaryDirectory() as folder:
            src = Path(folder)/'sample.docx'
            doc = Document(); doc.add_paragraph('文档'); doc.save(src)
            w = MainWindow()
            w.add_files([str(src), str(src), str(Path(folder)/'absent.docx')])
            self.assertEqual(w.filesList.count(), 1)
            w.lsType.setCurrentIndex(w.lsType.findData('exact')); w.lsValue.setValue(28)
            self.assertEqual(w.lsValue.value(),28)
            w.lsType.setCurrentIndex(w.lsType.findData('multiple')); w.lsValue.setValue(.5)
            self.assertEqual(w.lsValue.value(),.5)
            w._start()
            loop=QEventLoop(); w.worker.finished.connect(loop.quit); QTimer.singleShot(10000,loop.quit); loop.exec()
            self.assertFalse(w.worker.isRunning())
            self.app.processEvents()
            self.assertEqual(len(w.outputs),1)
            self.assertEqual('已完成', w.filesList.item(0).text(0))
            self.assertIn('正文', w.filesList.item(0).text(2))
            self.assertTrue(w.btnStart.isEnabled())
            self.assertFalse(w.btnCancel.isEnabled())
            w.close()

    def test_cancel_after_current_file(self):
        with tempfile.TemporaryDirectory() as folder:
            files=[]
            for i in range(3):
                path=Path(folder)/f'{i}.docx'; d=Document(); d.add_paragraph('test'); d.save(path); files.append(str(path))
            worker=FormatWorker(files, merge_preset({}), 'suffix', '', '_done')
            results=[]
            worker.outputReady.connect(lambda s,o: (results.append(o), worker.requestInterruption()), Qt.DirectConnection)
            worker.start(); self.assertTrue(worker.wait(10000))
            self.assertEqual(len(results),1)
            self.assertFalse((Path(folder)/'1_done.docx').exists())

    def test_dirty_tracking_and_ls_memory(self):
        w = MainWindow()
        w.presetList.setCurrentRow(0)
        self.assertFalse(w._is_dirty())
        self.assertFalse(w.btnReset.isEnabled())
        w.firstIndent.setValue(0)
        self.assertTrue(w._is_dirty())
        self.assertTrue(w.btnReset.isEnabled())
        w.btnReset.click()
        self.assertFalse(w._is_dirty())
        # 切换行距类型时各自记住数值
        w.lsType.setCurrentIndex(w.lsType.findData('exact')); w.lsValue.setValue(28)
        w.lsType.setCurrentIndex(w.lsType.findData('multiple'))
        self.assertEqual(w.lsValue.value(), 1.5)
        w.lsType.setCurrentIndex(w.lsType.findData('exact'))
        self.assertEqual(w.lsValue.value(), 28)
        w.close()

    def test_validation_blocks_bad_input(self):
        with tempfile.TemporaryDirectory() as folder:
            src = Path(folder)/'sample.docx'
            doc = Document(); doc.add_paragraph('文档'); doc.save(src)
            w = MainWindow()
            w.add_files([str(src)])
            w.pageSize.setCurrentIndex(w.pageSize.findData('B5'))  # 宽 18.2 cm
            w.margins['left'].setValue(9.9); w.margins['right'].setValue(9.9)
            self.assertIn('页边距', w._validate(w.collect_preset()))
            w.margins['left'].setValue(3); w.margins['right'].setValue(3)
            w.radioDir.setChecked(True); w.outdirEdit.setText('')
            self.assertIn('输出目录', w._validate(w.collect_preset()))
            w.radioSuffix.setChecked(True); w.suffixEdit.setText('a/b')
            self.assertIn('后缀', w._validate(w.collect_preset()))
            w.suffixEdit.setText('_ok')
            self.assertIsNone(w._validate(w.collect_preset()))
            w.close()

    def test_wps_keeps_wps_extension_and_falls_back(self):
        with tempfile.TemporaryDirectory() as folder:
            real = Path(folder)/'real.docx'
            doc = Document(); doc.add_paragraph('WPS 文档'); doc.save(real)
            wps = Path(folder)/'sample.wps'; wps.write_bytes(b'fake wps')
            def fake_to_docx(src, dst):
                shutil.copyfile(real, dst); return True
            def fake_to_wps(src, dst):
                shutil.copyfile(src, dst); return True
            outputs, done = [], []
            worker = FormatWorker([str(wps)], merge_preset({}), 'suffix', '', '_done')
            worker.outputReady.connect(lambda s, o: outputs.append(o))
            worker.fileDone.connect(lambda p, ok, msg: done.append((ok, msg)))
            with patch('app.gui.engine.convert_doc_to_docx', side_effect=fake_to_docx),                  patch('app.gui.engine.convert_docx_to_wps', side_effect=fake_to_wps):
                worker.run()
            self.assertEqual(Path(outputs[0]).name, 'sample_done.wps')
            self.assertTrue(done[0][0])
            self.assertEqual(Document(outputs[0]).paragraphs[0].text, 'WPS 文档')  # 假 WPS 只是复制
            self.assertFalse(list(Path(tempfile.gettempdir()).glob('docfmt_out_*.docx')))
            # WPS 不可用：退回 .docx，并在结果里说明
            outputs.clear(); done.clear()
            worker = FormatWorker([str(wps)], merge_preset({}), 'suffix', '', '_done')
            worker.outputReady.connect(lambda s, o: outputs.append(o))
            worker.fileDone.connect(lambda p, ok, msg: done.append((ok, msg)))
            with patch('app.gui.engine.convert_doc_to_docx', side_effect=fake_to_docx),                  patch('app.gui.engine.convert_docx_to_wps', return_value=False):
                worker.run()
            self.assertEqual(Path(outputs[0]).name, 'sample_done.docx')
            self.assertIn('已输出 .docx', done[0][1])

    def test_font_detection_and_user_install(self):
        import winreg
        test_key = r'Software\office-no1-test\Fonts'
        with tempfile.TemporaryDirectory() as folder,              patch.object(fonts, '_REG_FONTS', test_key),              patch.object(fonts, 'user_font_dir', return_value=folder):
            font = fonts.BUNDLED_FONTS[0]
            self.assertFalse(fonts.is_installed(font, registered=set(), families=set()))
            self.assertTrue(fonts.is_installed(font, registered={'fangsong_gb2312'}, families=set()))
            self.assertIsNone(fonts.install_font(font))
            self.assertTrue((Path(folder)/font['file']).is_file())
            self.assertIn(font['names'][0].lower(), fonts._registered_font_names())
            self.assertIsNone(fonts.install_font(font))  # 重复安装幂等
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, test_key, 0, winreg.KEY_ALL_ACCESS) as k:
                winreg.DeleteValue(k, f"{font['names'][0]} (TrueType)")
            winreg.DeleteKey(winreg.HKEY_CURRENT_USER, test_key)
            winreg.DeleteKey(winreg.HKEY_CURRENT_USER, r'Software\office-no1-test')

if __name__ == '__main__':
    unittest.main()
