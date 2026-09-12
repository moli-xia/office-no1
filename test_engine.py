"""引擎自动化测试：生成样例文档 → 套用全部内置预设 → 校验输出。"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
# Windows 控制台默认 GBK，✔/✘ 无法编码
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.engine import PAGE_SIZES_CM, apply_preset
from app.presets import BUILTIN_PRESETS

TEST_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_out")


def make_sample(path):
    doc = Document()
    doc.add_heading("测试文档一级标题", level=1)
    doc.add_heading("二级标题示例", level=2)
    p = doc.add_paragraph()
    r = p.add_run("这是一段正文测试文字，用于验证排版效果，包含中文与 English 混排。")
    r.font.name = "Calibri"  # 直接格式，验证运行级字体覆盖
    doc.add_paragraph("第二段正文，验证首行缩进与行距。")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "表格A"
    t.cell(1, 1).text = "表格B"
    doc.add_heading("三级标题示例", level=3)
    doc.add_heading("四级标题示例", level=4)
    doc.add_paragraph("结尾段落。")
    doc.save(path)


def cm(v):
    return getattr(v, "cm", None) or v / 360000.0


def check_output(out, preset, errors):
    doc = Document(out)
    page = preset["page"]
    sec = doc.sections[0]

    # 页面尺寸
    w, h = PAGE_SIZES_CM[page["size"]]
    if page["orientation"] == "landscape":
        w, h = h, w
    if abs(cm(sec.page_width) - w) > 0.05 or abs(cm(sec.page_height) - h) > 0.05:
        errors.append(f"[{preset['name']}] 页面尺寸不符: {cm(sec.page_width)}x{cm(sec.page_height)}")

    # 页边距
    for key, attr in (("top", "top_margin"), ("bottom", "bottom_margin"),
                      ("left", "left_margin"), ("right", "right_margin")):
        want = float(page[f"margin_{key}_cm"])
        got = cm(getattr(sec, attr))
        if abs(got - want) > 0.05:
            errors.append(f"[{preset['name']}] {attr} 期望 {want}cm 实际 {got:.3f}cm")

    # 文档网格必须被禁用（否则固定 / 倍数行距在 WPS 中会失真）
    grid = sec._sectPr.find(qn("w:docGrid"))
    if grid is None or grid.get(qn("w:type")) != "default" \
            or grid.get(qn("w:linePitch")) is not None or grid.get(qn("w:charSpace")) is not None:
        errors.append(f"[{preset['name']}] 文档网格未禁用: {None if grid is None else dict(grid.attrib)}")

    # Normal 样式字体与字号（Length 为 EMU，1 磅 = 12700 EMU）
    normal = doc.styles["Normal"]
    if abs(int(normal.font.size) - preset["body"]["size_pt"] * 12700) > 1:
        errors.append(f"[{preset['name']}] Normal 字号不符: {normal.font.size}")
    east = normal.element.rPr.rFonts.get(qn("w:eastAsia"))
    if east != preset["body"]["font_east"]:
        errors.append(f"[{preset['name']}] Normal 中文字体不符: {east}")

    # 正文段落：首行缩进(字符) + snapToGrid + 行距
    body_p = None
    for p in doc.paragraphs:
        if p.text.startswith("第二段正文"):
            body_p = p
            break
    if body_p is None:
        errors.append(f"[{preset['name']}] 未找到正文测试段")
    else:
        ind = body_p._p.pPr.find(qn("w:ind"))
        chars = float(preset["body"]["first_line_indent_chars"])
        if chars > 0 and (ind is None or ind.get(qn("w:firstLineChars")) != str(int(round(chars * 100)))):
            errors.append(f"[{preset['name']}] 首行缩进不符: {None if ind is None else ind.get(qn('w:firstLineChars'))}")
        snap = body_p._p.pPr.find(qn("w:snapToGrid"))
        if snap is None or snap.get(qn("w:val")) != "0":
            errors.append(f"[{preset['name']}] snapToGrid 应为 0")
        ls = body_p.paragraph_format.line_spacing
        want_ls = float(preset["body"]["line_spacing_value"])
        if preset["body"]["line_spacing_type"] == "multiple":
            if not (isinstance(ls, float) and abs(ls - want_ls) < 0.01):
                errors.append(f"[{preset['name']}] 行距不符: {ls!r}")
        else:
            if ls is None or abs(int(ls) - want_ls * 12700) > 1:
                errors.append(f"[{preset['name']}] 固定行距不符: {ls!r}")

    # 页码域
    pn = preset["page_number"]
    part = sec.footer if pn["position"].startswith("footer") else sec.header
    xml = part._element.xml if pn["enabled"] else ""
    if pn["enabled"] and "PAGE" not in xml:
        errors.append(f"[{preset['name']}] 页脚缺少 PAGE 域")
    if not pn["enabled"] and "PAGE" in xml:
        errors.append(f"[{preset['name']}] 页码应已移除但仍存在")
    if pn["enabled"] and pn["position"] == "footer_outside":
        # GB/T 9704：奇偶页不同，单页居右空一字、双页居左空一字，一字线
        if not doc.settings.odd_and_even_pages_header_footer:
            errors.append(f"[{preset['name']}] 未开启奇偶页不同")
        odd_p = sec.footer.paragraphs[0]._p
        even_p = sec.even_page_footer.paragraphs[0]._p
        if (str(odd_p.pPr.jc.val) != "RIGHT (2)" or odd_p.pPr.ind.get(qn("w:rightChars")) != "100"
                or str(even_p.pPr.jc.val) != "LEFT (0)" or even_p.pPr.ind.get(qn("w:leftChars")) != "100"):
            errors.append(f"[{preset['name']}] 奇偶页页码对齐 / 空一字不符")
        if "PAGE" not in sec.even_page_footer._element.xml:
            errors.append(f"[{preset['name']}] 偶数页页脚缺少 PAGE 域")
        if pn["format"] == "gb_dash" and "— " not in sec.footer.paragraphs[0].text:
            errors.append(f"[{preset['name']}] 页码一字线不符: {sec.footer.paragraphs[0].text!r}")
        # 页脚段落：单倍行距、不对齐网格
        fp = sec.footer.paragraphs[0]
        if fp.paragraph_format.line_spacing != 1.0 or fp._p.pPr.find(qn("w:snapToGrid")).get(qn("w:val")) != "0":
            errors.append(f"[{preset['name']}] 页脚段落行距 / 网格设置不符")
    elif pn["enabled"] and doc.settings.odd_and_even_pages_header_footer:
        errors.append(f"[{preset['name']}] 非公文位置不应开启奇偶页不同")

    # 一级标题格式（字体 + 加粗）
    h1 = None
    for p in doc.paragraphs:
        if p.text == "测试文档一级标题":
            h1 = p
            break
    if h1 is None:
        errors.append(f"[{preset['name']}] 未找到一级标题")
    else:
        hcfg = preset["headings"][0]
        rpr = h1._p.find(qn("w:r") ).find(qn("w:rPr")) if h1._p.find(qn("w:r")) is not None else None
        if rpr is None:
            errors.append(f"[{preset['name']}] 标题 run 无 rPr")
        else:
            got_east = rpr.rFonts.get(qn("w:eastAsia")) if rpr.rFonts is not None else None
            if got_east != hcfg["font_east"]:
                errors.append(f"[{preset['name']}] 一级标题中文字体不符: {got_east}")
            b = rpr.find(qn("w:b"))
            want_bold = "1" if hcfg["bold"] else "0"
            if b is None or b.get(qn("w:val")) != want_bold:
                errors.append(f"[{preset['name']}] 一级标题加粗不符")


def make_complex_sample(path):
    """多分节 + 自定义标题样式 + 编号列表 + 已有页眉，用于验证边界情况。"""
    from docx.enum.section import WD_SECTION
    from docx.oxml import OxmlElement

    doc = Document()
    # 基于 Heading 1 的自定义样式（常见模板做法）
    custom = doc.styles.add_style("章标题", 1)  # WD_STYLE_TYPE.PARAGRAPH
    custom.base_style = doc.styles["Heading 1"]
    doc.add_paragraph("自定义样式章标题", style=custom)
    doc.add_paragraph("普通正文段落。")
    lst = doc.add_paragraph("编号列表项", style="List Number")
    ind = lst.paragraph_format
    ind.left_indent = Cm(1.0)
    ind.first_line_indent = Cm(-0.5)
    doc.sections[0].header.paragraphs[0].text = "原有页眉"
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    sec2._sectPr.append(pg)
    doc.add_paragraph("第二节正文。")
    doc.save(path)


def check_complex(errors):
    import copy
    from app.presets import DEFAULT_PRESET

    sample = os.path.join(TEST_DIR, "样例_复杂.docx")
    make_complex_sample(sample)

    # 1) 启用页码：自定义标题样式被识别、列表缩进保留、第二节不重新编号
    preset = copy.deepcopy(DEFAULT_PRESET)
    preset["page_number"]["start_at"] = 3
    preset["page_number"]["hide_on_first_page"] = True
    out = os.path.join(TEST_DIR, "样例_复杂_页码.docx")
    res = apply_preset(sample, preset, out)
    if not res["ok"]:
        errors.append(f"[复杂] 处理失败: {res['error']}")
        return
    if res["stats"]["h1"] != 1:
        errors.append(f"[复杂] 基于 Heading 1 的自定义样式未被识别为一级标题: {res['stats']}")
    if res["stats"]["h4"] != 0 or res["stats"]["skipped"] != 0:
        errors.append(f"[复杂] 统计不符: {res['stats']}")
    doc = Document(out)
    lst = next(p for p in doc.paragraphs if p.text == "编号列表项")
    ind = lst._p.pPr.find(qn("w:ind"))
    if ind is None or ind.get(qn("w:hanging")) is None or ind.get(qn("w:firstLineChars")) is not None:
        errors.append("[复杂] 编号列表段落的悬挂缩进被改动")
    s1, s2 = doc.sections
    if s1._sectPr.find(qn("w:pgNumType")).get(qn("w:start")) != "3":
        errors.append("[复杂] 第一节起始页码未设置为 3")
    pg2 = s2._sectPr.find(qn("w:pgNumType"))
    if pg2 is not None and pg2.get(qn("w:start")) is not None:
        errors.append("[复杂] 第二节仍会重新开始编号")
    if not s1.different_first_page_header_footer or s2.different_first_page_header_footer:
        errors.append("[复杂] 首页不同 应只作用于第一节")
    if not s2.footer.is_linked_to_previous:
        errors.append("[复杂] 第二节页脚未链接到前一节")
    if "原有页眉" in s1.header._element.xml:
        errors.append("[复杂] 启用页码时页眉应被重建")

    # 2) 关闭页码且无页眉页脚文字：原文档页眉原样保留
    preset = copy.deepcopy(DEFAULT_PRESET)
    preset["page_number"]["enabled"] = False
    out = os.path.join(TEST_DIR, "样例_复杂_保留页眉.docx")
    res = apply_preset(sample, preset, out)
    if not res["ok"]:
        errors.append(f"[复杂] 处理失败: {res['error']}")
        return
    doc = Document(out)
    if "原有页眉" not in doc.sections[0].header._element.xml:
        errors.append("[复杂] 未启用页码时原有页眉应保留")
    normal_ind = doc.styles["Normal"].element.pPr.find(qn("w:ind")) if doc.styles["Normal"].element.pPr is not None else None
    if normal_ind is None or normal_ind.get(qn("w:firstLineChars")) != "200":
        errors.append("[复杂] Normal 样式首行缩进未写入")


def main():
    os.makedirs(TEST_DIR, exist_ok=True)
    sample = os.path.join(TEST_DIR, "样例.docx")
    make_sample(sample)

    errors = []
    for preset in BUILTIN_PRESETS:
        safe_name = preset["name"].replace("/", "-")  # Windows 文件名不能含 /
        out = os.path.join(TEST_DIR, f"样例_{safe_name}.docx")
        res = apply_preset(sample, preset, out)
        if not res["ok"]:
            errors.append(f"[{preset['name']}] 处理失败: {res['error']}")
            continue
        print(f"✔ {preset['name']:<28} -> {os.path.basename(out)}  ({res['message']})")
        check_output(out, preset, errors)

    check_complex(errors)
    print("✔ 复杂文档（多分节 / 自定义标题样式 / 列表 / 原有页眉）校验完成")

    if errors:
        print("\n发现 %d 个问题:" % len(errors))
        for e in errors:
            print("  ✘", e)
        sys.exit(1)
    print("\n全部 %d 个内置预设校验通过。" % len(BUILTIN_PRESETS))


if __name__ == "__main__":
    main()
